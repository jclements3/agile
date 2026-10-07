#!/usr/bin/env python3
"""
gitlab_agile_setup.py -- set up GitLab Premium/Ultimate for the tome / epoch / story / task
plan and the TODO/OPEN/DONE/WAIT/HOLD/PUNT/DROP work states, then load the SE backlog.

What it builds (each step is idempotent: existing items with the same title are reused):
  1. Work-item statuses   custom lifecycle with the 7 states and their colors (GitLab 18.2+)
  2. Labels               marks REDO / PASS / SYNC (olive), team::SE-*, and, with --label-states,
                          state::* labels for instances without custom statuses
  3. Iteration cadence    2-week sprints
  4. Milestones           one per epoch, named for the review that closes it
  5. Board                "Sprint board" with one list per state (label mode) or a plain board
  6. Issue templates      story.md and hardening.md committed to the project
  7. Backlog (optional)   tomes -> epics, epochs -> child epics, stories -> issues, tasks -> tasks,
                          read from se-backlog-staffing.xlsx, including each item's state and marks
  Export mode             --export FILE.xlsx writes GitLab back out in the same format, so a group
                          can be torn down and rebuilt (or moved) with --backlog FILE.xlsx

Usage (dry run is the default; nothing is written until you add --apply):
  export GITLAB_URL=https://gitlab.example.invalid
  export GITLAB_TOKEN=glpat-...            # scope: api; Owner/Maintainer of the top-level group
  python3 gitlab_agile_setup.py --group se-team --project se-team/se-backlog \
          --start 2026-10-12 --backlog se-backlog-staffing.xlsx            # preview
  python3 gitlab_agile_setup.py ... --apply                                # do it
  python3 gitlab_agile_setup.py --group se-team --project se-team/se-backlog \
          --export snapshot.xlsx --template se-backlog-staffing.xlsx       # GitLab -> xlsx

Requirements: Python 3.8+, stdlib only. openpyxl is needed only for --backlog and --export.
"""
import argparse, datetime as dt, json, os, ssl, sys, time, urllib.parse, urllib.request

# ---------------------------------------------------------------- the scheme
STATES = [  # name, hex color, GitLab status category, description
    ("TODO", "#808082", "TO_DO",       "On the team backlog, not started"),
    ("OPEN", "#27235d", "IN_PROGRESS", "Begun, in the sprint"),
    ("WAIT", "#d6292e", "IN_PROGRESS", "Blocked by something outside the team"),
    ("HOLD", "#c98500", "IN_PROGRESS", "Interrupted by higher priority (hold / resume)"),
    ("PUNT", "#d6292e", "TRIAGE",      "Too hard as written, back to TODO to be re-cut"),
    ("DONE", "#0ca30c", "DONE",        "Completed and accepted at the demo"),
    ("DROP", "#555555", "CANCELLED",   "Should not be done"),
]
DEFAULT_OPEN, DEFAULT_CLOSED, DEFAULT_DUPLICATE = "TODO", "DONE", "DROP"
MARKS = [("REDO", "The demo showed it was wrong, back in this sprint"),
         ("PASS", "Handed to another team (pass ID Team)"),
         ("SYNC", "Coordinated across teams, with a shared DONE")]
MARK_COLOR = "#6b7436"
TEAMS = ["SE-ARCH", "SE-REQ", "SE-INT", "SE-PERF", "SE-VV", "SE-SPEC"]
TEAM_COLOR = "#2f5d8a"
SPRINT_WEEKS, SPRINTS_PER_EPOCH = 2, 13
EPOCH_REVIEWS = {"T1.E1": "SRR", "T1.E2": "SFR", "T1.E3": "PDR", "T1.E4": "CDR",
                 "T1.E5": "TRR", "T1.E6": "SVR/FRR"}

TEMPLATES = {
".gitlab/issue_templates/story.md": """## Story
<!-- One sprint (2 weeks). Split anything bigger: set status PUNT and re-cut. -->

**Done when:**
- [ ]

**Model link:** <!-- requirement / verification case / model element this story changes -->

## Tasks
<!-- Add child tasks, each under one week. -->

## Definition of done
- [ ] Acceptance criteria demonstrated at the sprint demo
- [ ] Model changes merged, peer-reviewed, pipeline green
- [ ] Trace links updated (derive / satisfy / verify)

/label ~"team::SE-"
""",
".gitlab/issue_templates/hardening.md": """## Sprint 13 hardening
- [ ] Integration fixes from the epoch closed or carried as WAIT with owner
- [ ] Segment configurations frozen and tagged
- [ ] End-to-end run in the integration lab
- [ ] Model baseline tagged; review package generated from the model
- [ ] Review held; actions logged as TODO in the next epoch
""",
}

# ---------------------------------------------------------------- HTTP helpers
class GL:
    def __init__(self, url, token, apply, insecure=False):
        self.url, self.token, self.apply = url.rstrip("/"), token, apply
        self.ctx = ssl._create_unverified_context() if insecure else None

    def _req(self, method, path, data=None, query=None):
        u = f"{self.url}/api/v4{path}"
        if query: u += "?" + urllib.parse.urlencode(query, doseq=True)
        body = json.dumps(data).encode() if data is not None else None
        r = urllib.request.Request(u, data=body, method=method,
                                   headers={"PRIVATE-TOKEN": self.token, "Content-Type": "application/json"})
        for attempt in range(4):
            try:
                with urllib.request.urlopen(r, context=self.ctx) as resp:
                    txt = resp.read().decode()
                    return json.loads(txt) if txt else None, resp.headers
            except urllib.error.HTTPError as e:
                if e.code == 429 and attempt < 3:
                    time.sleep(2 ** attempt); continue
                raise RuntimeError(f"{method} {path} -> {e.code}: {e.read().decode()[:300]}")

    def get_all(self, path, query=None):
        q, out, page = dict(query or {}), [], 1
        while True:
            q.update(per_page=100, page=page)
            data, hdr = self._req("GET", path, query=q)
            out += data or []
            nxt = hdr.get("X-Next-Page")
            if not nxt: return out
            page = int(nxt)

    def write(self, method, path, data, what):
        print(f"  {'+' if self.apply else '~'} {what}")
        if not self.apply: return {"id": 0, "iid": 0, "web_url": "(dry run)"}
        return self._req(method, path, data)[0]

    def gql(self, query, variables=None, mutate=False, what=""):
        if mutate:
            print(f"  {'+' if self.apply else '~'} {what}")
            if not self.apply: return {}
        r = urllib.request.Request(f"{self.url}/api/graphql",
                                   data=json.dumps({"query": query, "variables": variables or {}}).encode(),
                                   headers={"Authorization": f"Bearer {self.token}", "Content-Type": "application/json"})
        with urllib.request.urlopen(r, context=self.ctx) as resp:
            out = json.loads(resp.read().decode())
        if out.get("errors"): raise RuntimeError(json.dumps(out["errors"])[:400])
        if mutate:  # mutations report validation failures in the payload, not as GraphQL errors
            for payload in (out.get("data") or {}).values():
                if isinstance(payload, dict) and payload.get("errors"):
                    raise RuntimeError(f"{what}: {payload['errors']}")
        return out["data"]

def enc(p): return urllib.parse.quote(p, safe="")

# ---------------------------------------------------------------- 1. statuses
def setup_statuses(gl, group):
    print("\n[1] Work-item statuses (custom lifecycle)")
    try:
        d = gl.gql("""query($p:ID!){ namespace(fullPath:$p){ lifecycles{ nodes{ id name
                      statuses{ id name } workItemTypes{ name } } } } }""", {"p": group})
        nodes = (d["namespace"] or {}).get("lifecycles", {}).get("nodes", [])
    except Exception as e:
        print(f"  ! could not read lifecycles ({e}).\n"
              "    Custom statuses need GitLab 18.2+ Premium/Ultimate. Set them by hand under\n"
              "    Group > Settings > Issues > Statuses, or rerun with --label-states.")
        return None
    target = next((n for n in nodes if any(t["name"] == "Issue" for t in n.get("workItemTypes", []))), None) \
             or (nodes[0] if nodes else None)
    if not target:
        print("  ! no lifecycle found to update; set statuses in the UI."); return None
    existing = {s["name"]: s["id"] for s in target["statuses"]}
    statuses = []
    for name, color, cat, desc in STATES:
        s = {"name": name, "color": color, "category": cat, "description": desc}
        if name in existing: s["id"] = existing[name]
        statuses.append(s)
    names = [s[0] for s in STATES]
    d = gl.gql("""mutation($input:LifecycleUpdateInput!){ lifecycleUpdate(input:$input){
                lifecycle{ id name statuses{ id name } } errors } }""",
           {"input": {"namespacePath": group, "id": target["id"], "statuses": statuses,
                      "defaultOpenStatusIndex": names.index(DEFAULT_OPEN),
                      "defaultClosedStatusIndex": names.index(DEFAULT_CLOSED),
                      "defaultDuplicateStatusIndex": names.index(DEFAULT_DUPLICATE)}},
           mutate=True, what=f"lifecycle '{target['name']}' -> {', '.join(names)}")
    print("    note: statuses not in this list are removed unless in use; in-use ones must be"
          " remapped in the UI first.")
    lc = ((d or {}).get("lifecycleUpdate") or {}).get("lifecycle") or {}
    return {x["name"]: x["id"] for x in lc.get("statuses", [])}

# ---------------------------------------------------------------- 2. labels
def setup_labels(gl, group, label_states):
    print("\n[2] Labels")
    have = {l["name"]: l for l in gl.get_all(f"/groups/{enc(group)}/labels", {"include_ancestor_groups": False})}
    want = [(m, MARK_COLOR, d) for m, d in MARKS] + [(f"team::{t}", TEAM_COLOR, f"Owned by {t}") for t in TEAMS]
    if label_states:
        want += [(f"state::{n}", c, d) for n, c, _, d in STATES]
    ids = {}
    for name, color, desc in want:
        if name in have:
            ids[name] = have[name]["id"]; print(f"  = {name}"); continue
        r = gl.write("POST", f"/groups/{enc(group)}/labels",
                     {"name": name, "color": color, "description": desc}, f"label {name} {color}")
        ids[name] = r["id"]
    return ids

# ---------------------------------------------------------------- 3. cadence
def setup_cadence(gl, group, start):
    print("\n[3] Iteration cadence")
    d = gl.gql("""query($p:ID!){ group(fullPath:$p){ iterationCadences{ nodes{ id title } } } }""", {"p": group})
    title = "SE sprints (2 weeks)"
    if any(n["title"] == title for n in d["group"]["iterationCadences"]["nodes"]):
        print(f"  = {title}"); return
    gl.gql("""mutation($i:IterationCadenceCreateInput!){ iterationCadenceCreate(input:$i){
                iterationCadence{ id } errors } }""",
           {"i": {"groupPath": group, "title": title, "automatic": True, "active": True,
                  "startDate": start.isoformat(), "durationInWeeks": SPRINT_WEEKS,
                  "iterationsInAdvance": SPRINTS_PER_EPOCH, "rollOver": True,
                  "description": "Two-week sprints; 13 per epoch; sprint 13 is hardening."}},
           mutate=True, what=f"cadence '{title}' from {start}, {SPRINT_WEEKS}-week, roll-over on")

# ---------------------------------------------------------------- 4. milestones
def epoch_dates(start, k):
    s = start + dt.timedelta(weeks=SPRINT_WEEKS * SPRINTS_PER_EPOCH * k)
    return s, s + dt.timedelta(weeks=SPRINT_WEEKS * SPRINTS_PER_EPOCH) - dt.timedelta(days=1)

def setup_milestones(gl, group, start, epochs):
    print("\n[4] Milestones (one per epoch)")
    have = {m["title"]: m["id"] for m in gl.get_all(f"/groups/{enc(group)}/milestones")}
    ids = {}
    for k, (eid, title) in enumerate(epochs):
        s, e = epoch_dates(start, k)
        name = f"{eid} {EPOCH_REVIEWS.get(eid, 'baseline')}"
        if name in have:
            ids[eid] = have[name]; print(f"  = {name}"); continue
        r = gl.write("POST", f"/groups/{enc(group)}/milestones",
                     {"title": name, "description": title, "start_date": s.isoformat(), "due_date": e.isoformat()},
                     f"milestone {name}  {s} .. {e}")
        ids[eid] = r["id"]
    return ids

# ---------------------------------------------------------------- 5. board
def setup_board(gl, project, label_ids, label_states):
    print("\n[5] Board")
    boards = gl.get_all(f"/projects/{enc(project)}/boards")
    b = next((x for x in boards if x["name"] == "Sprint board"), None)
    if not b:
        b = gl.write("POST", f"/projects/{enc(project)}/boards", {"name": "Sprint board"}, "board 'Sprint board'")
    else:
        print("  = Sprint board")
    if label_states:
        have = {l["label"]["name"] for l in (b.get("lists") or []) if l.get("label")}
        for n, *_ in STATES:
            if n in ("TODO", "DONE", "DROP"): continue   # Open/Closed columns already cover these
            name = f"state::{n}"
            if name in have: continue
            gl.write("POST", f"/projects/{enc(project)}/boards/{b['id']}/lists",
                     {"label_id": label_ids.get(name)}, f"board list {name}")
    else:
        print("  i  With custom statuses, add one list per status in the board UI\n"
              "     (board > Create list > Status) if your version offers status lists;\n"
              "     otherwise filter the board by Status. Scope the board to the current iteration.")

# ---------------------------------------------------------------- 6. templates
def setup_templates(gl, project, branch):
    print("\n[6] Issue templates")
    for path, content in TEMPLATES.items():
        try:
            gl._req("GET", f"/projects/{enc(project)}/repository/files/{enc(path)}", query={"ref": branch})
            print(f"  = {path}"); continue
        except RuntimeError:
            pass
        gl.write("POST", f"/projects/{enc(project)}/repository/files/{enc(path)}",
                 {"branch": branch, "content": content, "commit_message": f"Add issue template {path}"},
                 f"commit {path} on {branch}")

# ---------------------------------------------------------------- 7. backlog import
STATE_NAMES = [s[0] for s in STATES]
MARK_NAMES = [m[0] for m in MARKS]
HOURS_PER_DAY = 8          # GitLab default: 1d = 8h

def read_backlog(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Backlog"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    # Story days are SUMIFS formulas in the planning workbook; if the file was saved without cached
    # values, rebuild them from the tasks so issue weights are never silently zero.
    for r in rows:
        if r["Level"] == "Story" and not isinstance(r.get("Days"), (int, float)):
            r["Days"] = sum(t.get("Days") or 0 for t in rows if t["Level"] == "Task" and t["Parent"] == r["ID"]
                            and isinstance(t.get("Days"), (int, float))) or None
    return rows

def row_state(r):
    v = str(r.get("Status") or "").strip().upper()
    return v if v in STATE_NAMES else "TODO"

def row_marks(r):
    return [m for m in str(r.get("Marks") or "").replace(",", " ").upper().split() if m in MARK_NAMES]

class StateSetter:
    """Sets an item's state: custom status when available, else state::* label plus close."""
    def __init__(self, gl, project, status_ids, label_states):
        self.gl, self.project, self.ids, self.labels = gl, project, status_ids or {}, label_states
        self.warned = False

    def issue(self, iid, issue_id, state):
        if state == "TODO" or not self.gl.apply: return
        if not self.labels and self._status(f"gid://gitlab/WorkItem/{issue_id}", state): return
        if state in ("DONE", "DROP"):
            self.gl._req("PUT", f"/projects/{enc(self.project)}/issues/{iid}", {"state_event": "close"})

    def task(self, gid, state):
        if state == "TODO" or not self.gl.apply: return
        if not self.labels and self._status(gid, state): return
        if state in ("DONE", "DROP"):
            self.gl.gql("""mutation($i:WorkItemUpdateInput!){ workItemUpdate(input:$i){ errors } }""",
                        {"i": {"id": gid, "stateEvent": "CLOSE"}}, mutate=True, what=f"close task {gid}")

    def _status(self, gid, state):
        sid = self.ids.get(state)
        if not sid: return False
        try:
            self.gl.gql("""mutation($i:WorkItemUpdateInput!){ workItemUpdate(input:$i){ errors } }""",
                        {"i": {"id": gid, "statusWidget": {"status": sid}}}, mutate=True,
                        what=f"  status {state}")
            return True
        except Exception as e:
            if not self.warned:
                print(f"  ! could not set status ({e}); falling back to close-only"); self.warned = True
            return False

def load_backlog(gl, group, project, rows, milestones, start, status_ids, label_states):
    print("\n[7] Backlog import")
    epics = {e["title"]: e for e in gl.get_all(f"/groups/{enc(group)}/epics")}
    issues = {i["title"]: i for i in gl.get_all(f"/projects/{enc(project)}/issues", {"scope": "all"})}
    setter = StateSetter(gl, project, status_ids, label_states)
    task_type = None
    if gl.apply:
        d = gl.gql("""query($p:ID!){ project(fullPath:$p){ id workItemTypes(name:TASK){ nodes{ id } } } }""",
                   {"p": project})
        task_type = d["project"]["workItemTypes"]["nodes"][0]["id"]
    ref, made, epoch_index = {}, {"epic": 0, "issue": 0, "task": 0}, {}
    for r in rows:
        lvl, rid, title = r["Level"], r["ID"], f'{r["ID"]} {r["Title"]}'
        state, marks = row_state(r), row_marks(r)
        tag = f"  [{state}{' ' + ' '.join(marks) if marks else ''}]"
        if lvl in ("Tome", "Epoch"):
            if lvl == "Epoch": epoch_index[rid] = len(epoch_index)
            if title in epics:
                ref[rid] = epics[title]; print(f"  = epic {title}"); continue
            body = {"title": title, "description": f"**Done when:** {r.get('Done when') or ''}\n\n**State:** {state}"}
            if lvl == "Epoch" and r.get("Parent") in ref:
                body["parent_id"] = ref[r["Parent"]]["id"]
                s, e = epoch_dates(start, epoch_index[rid])
                body.update(start_date_is_fixed=True, start_date_fixed=s.isoformat(),
                            due_date_is_fixed=True, due_date_fixed=e.isoformat())
            if state in ("DONE", "DROP"): body["state_event"] = "close"
            ref[rid] = gl.write("POST", f"/groups/{enc(group)}/epics", body, f"epic {title}")
            made["epic"] += 1
        elif lvl == "Story":
            if title in issues:
                ref[rid] = issues[title]; print(f"  = issue {title}"); continue
            k = epoch_index.get(r["Epoch"], 0)
            sprint = int(r["Sprint"]) if str(r.get("Sprint")).isdigit() else SPRINTS_PER_EPOCH
            due = epoch_dates(start, k)[0] + dt.timedelta(weeks=SPRINT_WEEKS * sprint) - dt.timedelta(days=1)
            desc = (f"**Done when:** {r.get('Done when') or ''}\n\n**Model link:** {r.get('Model link') or ''}\n\n"
                    f"Planned sprint {sprint} of {r['Epoch']}. Assign the iteration at sprint planning.")
            labels = [f"team::{r['Team']}"] + marks + ([f"state::{state}"] if label_states else [])
            body = {"title": title, "description": desc, "labels": ",".join(labels),
                    "weight": int(r.get("Days") or 0), "due_date": due.isoformat(),
                    "milestone_id": milestones.get(r["Epoch"])}
            if r["Epoch"] in ref: body["epic_id"] = ref[r["Epoch"]]["id"]
            ref[rid] = gl.write("POST", f"/projects/{enc(project)}/issues", body,
                                f"issue {title}  ({r['Team']}, sprint {sprint}, {r.get('Days')} d){tag}")
            setter.issue(ref[rid].get("iid"), ref[rid].get("id"), state)
            made["issue"] += 1
        elif lvl == "Task":
            parent = ref.get(r["Parent"])
            what = f"task {title}  ({r.get('Days')} d){tag}"
            if not gl.apply or not parent:
                print(f"  ~ {what}"); made["task"] += 1; continue
            d = gl.gql("""mutation($i:WorkItemCreateInput!){ workItemCreate(input:$i){ workItem{ id } errors } }""",
                       {"i": {"projectPath": project, "title": title, "workItemTypeId": task_type,
                              "hierarchyWidget": {"parentId": f"gid://gitlab/WorkItem/{parent['id']}"},
                              "timeTrackingWidget": {"timeEstimate": f"{int(r.get('Days') or 0)}d"}}},
                       mutate=True, what=what)
            setter.task(d["workItemCreate"]["workItem"]["id"], state)
            made["task"] += 1
    print(f"  -> {made['epic']} epics, {made['issue']} issues, {made['task']} tasks "
          f"{'created' if gl.apply else 'would be created'}")

# ---------------------------------------------------------------- 8. export
EXPORT_QUERY = """query($p:ID!,$after:String){ project(fullPath:$p){
  workItems(first:100, after:$after, types:[ISSUE, TASK]){ pageInfo{ hasNextPage endCursor }
    nodes{ id iid title state webUrl workItemType{ name }
      widgets{ __typename
        STATUS_FRAGMENT
        ... on WorkItemWidgetHierarchy{ parent{ iid workItemType{ name } } }
        ... on WorkItemWidgetTimeTracking{ timeEstimate totalTimeSpent }
        ... on WorkItemWidgetLabels{ labels{ nodes{ title } } }
        ... on WorkItemWidgetAssignees{ assignees{ nodes{ username } } } } } } } }"""
STATUS_FRAGMENT = "... on WorkItemWidgetStatus{ status{ name } }"

def fetch_work_items(gl, project):
    for frag in (STATUS_FRAGMENT, ""):          # retry without statuses on older instances
        try:
            out, after = [], None
            while True:
                d = gl.gql(EXPORT_QUERY.replace("STATUS_FRAGMENT", frag), {"p": project, "after": after})
                wi = d["project"]["workItems"]
                out += wi["nodes"]
                if not wi["pageInfo"]["hasNextPage"]: return out
                after = wi["pageInfo"]["endCursor"]
        except RuntimeError as e:
            if frag and "WorkItemWidgetStatus" in str(e): continue
            raise

def split_id(title, fallback):
    head, _, rest = (title or "").partition(" ")
    if head[:1] == "T" and head[1:2].isdigit():     # backlog IDs: T1, T1.E2, T1.E2.S03, T1.E2.S03.K1
        return head, rest
    return fallback, title

def widget(node, typename):
    return next((w for w in node.get("widgets", []) if w.get("__typename") == typename), {}) or {}

def item_state(status_name, labels, closed):
    if status_name and status_name.upper() in STATE_NAMES: return status_name.upper()
    for l in labels:
        if l.startswith("state::") and l[7:] in STATE_NAMES: return l[7:]
    return "DONE" if closed else "TODO"

def grab(desc, key):
    for line in (desc or "").splitlines():
        if line.startswith(f"**{key}:**"): return line.split(":**", 1)[1].strip()
    return ""

def export_backlog(gl, group, project, out_path, template):
    print(f"\n[export] {group} / {project} -> {out_path}")
    epics = gl.get_all(f"/groups/{enc(group)}/epics", {"include_descendant_groups": False})
    issues = gl.get_all(f"/projects/{enc(project)}/issues", {"scope": "all"})
    wis = fetch_work_items(gl, project)
    wi_by_iid = {str(w["iid"]): w for w in wis}
    by_id = {e["id"]: e for e in epics}

    rows = []
    eid = {e["id"]: split_id(e["title"], f"EPIC-{e['iid']}")[0] for e in epics}
    def epic_row(e, lvl):
        rid, title = split_id(e["title"], eid[e["id"]])
        parent = eid.get(e.get("parent_id")) if e.get("parent_id") else ""
        return {"ID": rid, "Level": lvl, "Parent": parent or "", "Title": title, "Team": "",
                "Epoch": rid if lvl == "Epoch" else "", "Sprint": "1-13" if lvl == "Epoch" else "",
                "Days": None, "Done when": grab(e.get("description"), "Done when"), "Model link": "",
                "Status": (grab(e.get("description"), "State").upper() if grab(e.get("description"), "State").upper()
                           in STATE_NAMES else ("DONE" if e.get("state") == "closed" else "TODO")),
                "Marks": "", "Iteration": "", "Assignee": "", "Spent (d)": None,
                "GitLab": e.get("web_url"), "_gl_id": e["id"]}
    tomes = [e for e in epics if not e.get("parent_id") or e["parent_id"] not in by_id]
    issues_by_epic, unplanned = {}, []
    for i in issues:
        if i.get("issue_type", "issue") != "issue": continue
        ep = (i.get("epic") or {}).get("id") or i.get("epic_id")
        (issues_by_epic.setdefault(ep, []) if ep in by_id else unplanned).append(i)
    tasks_by_parent = {}
    for w in wis:
        if (w.get("workItemType") or {}).get("name") != "Task": continue
        p = (widget(w, "WorkItemWidgetHierarchy").get("parent") or {}).get("iid")
        tasks_by_parent.setdefault(str(p), []).append(w)

    def story_rows(i, epoch_id):
        w = wi_by_iid.get(str(i["iid"]), {})
        labels = i.get("labels") or []
        st = item_state((widget(w, "WorkItemWidgetStatus").get("status") or {}).get("name"),
                        labels, i.get("state") == "closed")
        rid, title = split_id(i["title"], f"ISSUE-{i['iid']}")
        team = next((l.split("::", 1)[1] for l in labels if l.startswith("team::")), "")
        sprint, desc = "", i.get("description") or ""
        if "Planned sprint " in desc:
            sprint = desc.split("Planned sprint ", 1)[1].split()[0]
        ts = i.get("time_stats") or {}
        rows.append({"ID": rid, "Level": "Story", "Parent": epoch_id, "Title": title, "Team": team,
                     "Epoch": epoch_id, "Sprint": int(sprint) if str(sprint).isdigit() else sprint,
                     "Days": i.get("weight"), "Done when": grab(desc, "Done when"),
                     "Model link": grab(desc, "Model link"), "Status": st,
                     "Marks": " ".join(m for m in MARK_NAMES if m in labels),
                     "Iteration": (i.get("iteration") or {}).get("title") or "",
                     "Assignee": ", ".join(a["username"] for a in i.get("assignees") or []),
                     "Spent (d)": round((ts.get("total_time_spent") or 0) / 3600 / HOURS_PER_DAY, 1) or None,
                     "GitLab": i.get("web_url")})
        for k, t in enumerate(sorted(tasks_by_parent.get(str(i["iid"]), []), key=lambda t: t["title"]), 1):
            tt = widget(t, "WorkItemWidgetTimeTracking")
            tl = [n["title"] for n in (widget(t, "WorkItemWidgetLabels").get("labels") or {}).get("nodes", [])]
            tid, ttitle = split_id(t["title"], f"{rid}.K{k}")
            rows.append({"ID": tid, "Level": "Task", "Parent": rid, "Title": ttitle, "Team": team,
                         "Epoch": epoch_id, "Sprint": int(sprint) if str(sprint).isdigit() else sprint,
                         "Days": round((tt.get("timeEstimate") or 0) / 3600 / HOURS_PER_DAY, 1) or None,
                         "Done when": "", "Model link": "",
                         "Status": item_state((widget(t, "WorkItemWidgetStatus").get("status") or {}).get("name"),
                                              tl, t.get("state") == "CLOSED"),
                         "Marks": " ".join(m for m in MARK_NAMES if m in tl), "Iteration": "",
                         "Assignee": ", ".join(a["username"] for a in
                                               (widget(t, "WorkItemWidgetAssignees").get("assignees") or {}).get("nodes", [])),
                         "Spent (d)": round((tt.get("totalTimeSpent") or 0) / 3600 / HOURS_PER_DAY, 1) or None,
                         "GitLab": t.get("webUrl")})

    for tome in sorted(tomes, key=lambda e: e["title"]):
        rows.append(epic_row(tome, "Tome"))
        for ep in sorted([e for e in epics if e.get("parent_id") == tome["id"]], key=lambda e: e["title"]):
            rows.append(epic_row(ep, "Epoch"))
            for i in sorted(issues_by_epic.get(ep["id"], []), key=lambda i: i["title"]):
                story_rows(i, eid[ep["id"]])
        for i in sorted(issues_by_epic.get(tome["id"], []), key=lambda i: i["title"]):
            story_rows(i, eid[tome["id"]])
    if unplanned:
        rows.append({"ID": "UNPLANNED", "Level": "Epoch", "Parent": "", "Title": "Issues with no epic",
                     "Epoch": "UNPLANNED", "Status": "TODO"})
        for i in sorted(unplanned, key=lambda i: i["title"]): story_rows(i, "UNPLANNED")
    write_workbook(rows, out_path, template)
    n = {lvl: sum(1 for r in rows if r["Level"] == lvl) for lvl in ("Tome", "Epoch", "Story", "Task")}
    print(f"  -> {n['Tome']} tomes, {n['Epoch']} epochs, {n['Story']} stories, {n['Task']} tasks written")

COLS = ["ID", "Level", "Parent", "Title", "Team", "Epoch", "Sprint", "Days", "Done when", "Model link",
        "Status", "Marks", "Iteration", "Assignee", "Spent (d)", "GitLab"]

def write_workbook(rows, out_path, template):
    """Write the Backlog sheet in the importer's format. If a template workbook is given, its other
    sheets (Teams, Staffing, Profile, Capacity) are kept, so capacity formulas keep working."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = load_workbook(template) if template else Workbook()
    if "Backlog" in wb.sheetnames:
        idx = wb.sheetnames.index("Backlog"); del wb["Backlog"]; ws = wb.create_sheet("Backlog", idx)
    else:
        ws = wb.active; ws.title = "Backlog"
    F = "Arial"
    color = {n: c.lstrip("#") for n, c, *_ in STATES}
    fill = {"Tome": "D9D2F5", "Epoch": "FBEFD0", "Story": "FCE4DA"}
    for c, h in enumerate(COLS, 1):
        x = ws.cell(1, c, h); x.font = Font(name=F, bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", fgColor="2F3B4A")
    N = len(rows) + 1
    for r, row in enumerate(rows, 2):
        lvl = row["Level"]
        for c, h in enumerate(COLS, 1):
            v = row.get(h)
            if h == "Days" and lvl in ("Tome", "Epoch"):
                v = f"=SUMIFS($H$2:$H${N},$C$2:$C${N},A{r})"
            x = ws.cell(r, c, v if v not in ("",) else None)
            x.font = Font(name=F, bold=lvl in ("Tome", "Epoch"))
            if lvl in fill and h != "Status": x.fill = PatternFill("solid", fgColor=fill[lvl])
            if h == "Title": x.alignment = Alignment(indent={"Tome": 0, "Epoch": 1, "Story": 2, "Task": 3}[lvl])
            if h == "Status" and v in color:
                x.fill = PatternFill("solid", fgColor=color[v]); x.font = Font(name=F, bold=True, color="FFFFFF")
            if h == "Marks" and v:
                x.font = Font(name=F, bold=True, color=MARK_COLOR.lstrip("#"))
        ws.row_dimensions[r].outline_level = {"Tome": 0, "Epoch": 1, "Story": 2, "Task": 3}[lvl]
    for c, w in zip(range(1, len(COLS) + 1), [16, 8, 13, 50, 10, 10, 7, 7, 45, 26, 8, 12, 18, 14, 9, 40]):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = w
    ws.freeze_panes = "E2"; ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLS)).column_letter}{N}"
    ws.sheet_properties.outlinePr.summaryBelow = False
    wb.save(out_path)

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--group", required=True, help="top-level group path, e.g. se-team")
    ap.add_argument("--project", required=True, help="project path for issues/board, e.g. se-team/se-backlog")
    ap.add_argument("--start", help="first sprint start date YYYY-MM-DD (a Monday); needed for setup")
    ap.add_argument("--backlog", help="xlsx to import epics/issues/tasks from (states and marks included)")
    ap.add_argument("--export", metavar="XLSX", help="export the GitLab backlog to this xlsx and exit")
    ap.add_argument("--template", help="workbook whose Teams/Staffing/Capacity sheets are kept on export")
    ap.add_argument("--branch", default="main", help="branch for issue templates")
    ap.add_argument("--label-states", action="store_true",
                    help="use state::* scoped labels instead of custom statuses (pre-18.2 or Free)")
    ap.add_argument("--skip", default="", help="comma list of steps to skip: statuses,labels,cadence,"
                                               "milestones,board,templates,backlog")
    ap.add_argument("--apply", action="store_true", help="write changes (default is a dry run)")
    ap.add_argument("--insecure", action="store_true", help="skip TLS verification (self-signed lab servers)")
    a = ap.parse_args()
    url, token = os.environ.get("GITLAB_URL"), os.environ.get("GITLAB_TOKEN")
    if not url or not token: sys.exit("Set GITLAB_URL and GITLAB_TOKEN.")
    gl = GL(url, token, a.apply, a.insecure)

    if a.export:                                  # read-only; no --apply needed
        export_backlog(gl, a.group, a.project, a.export, a.template)
        return
    if not a.start: sys.exit("--start is required for setup/import.")
    start = dt.date.fromisoformat(a.start)
    skip = set(filter(None, a.skip.split(",")))
    print(f"{'APPLY' if a.apply else 'DRY RUN'} against {url}  group={a.group}  project={a.project}")
    rows = read_backlog(a.backlog) if a.backlog else []
    epochs = [(r["ID"], r["Title"]) for r in rows if r["Level"] == "Epoch" and str(r["ID"]).startswith("T1.")] \
             or [(f"T1.E{k}", f"Epoch {k}") for k in range(1, 7)]

    label_states, status_ids = a.label_states, {}
    if "statuses" not in skip and not label_states:
        status_ids = setup_statuses(gl, a.group)
        if status_ids is None: label_states = True; status_ids = {}; print("  -> falling back to state::* labels")
    label_ids = setup_labels(gl, a.group, label_states) if "labels" not in skip else {}
    if "cadence" not in skip: setup_cadence(gl, a.group, start)
    ms = setup_milestones(gl, a.group, start, epochs) if "milestones" not in skip else {}
    if "board" not in skip: setup_board(gl, a.project, label_ids, label_states)
    if "templates" not in skip: setup_templates(gl, a.project, a.branch)
    if rows and "backlog" not in skip:
        load_backlog(gl, a.group, a.project, rows, ms, start, status_ids, label_states)
    print("\nDone." if a.apply else "\nDry run complete. Rerun with --apply to make these changes.")

if __name__ == "__main__":
    main()
